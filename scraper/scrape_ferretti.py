#!/usr/bin/env python3
"""Scraper catalogo Ferretti Yachts -> Excel (foglio Macchine) + immagini.

Fonte: https://www.ferretti-yachts.com/it-it/ (pagine serie -> pagine modello, HTML statico).
Solo libreria standard + openpyxl. Pausa tra richieste, nessun login, nessun path vietato da robots.txt.
Prezzo = 0 (da compilare a mano / Excel listino).
Uso: python3 scrape_ferretti.py [--out FILE.xlsx] [--delay 1.5] [--images DIR]
"""
import argparse, html as H, os, re, sys, time, urllib.request
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = "https://www.ferretti-yachts.com"
SERIE = {"Flybridge": "Gamma flybridge", "Infynito": "Infynito"}
UA = "Mozilla/5.0 (compatible; CatalogBot/1.0)"

COLONNE = ["Codice", "Famiglia", "Categoria", "Sottocategoria", "Nome", "Prezzo",
           "LOA_m", "Larghezza_m", "Pescaggio_m", "Lunghezza_scafo_m", "Cabine",
           "Motore_Standard", "Vel_Max_Standard_kn", "Vel_Crociera_Standard_kn",
           "Motore_Optional", "Vel_Max_Optional_kn", "Vel_Crociera_Optional_kn",
           "Dislocamento_Vuoto_kg", "Dislocamento_Carico_kg", "Carburante_l", "Acqua_l", "Persone",
           "Descrizione", "Descrizione_EN", "Storia_IT", "Storia_EN", "URL", "Immagine_URL", "Immagine_File"]


def get(url, binary=False):
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    for t in range(3):
        try:
            with urllib.request.urlopen(req, timeout=40) as r:
                d = r.read()
                return d if binary else d.decode("utf-8", "ignore")
        except Exception:
            if t == 2:
                raise
            time.sleep(2 * (t + 1))


def righe(page):
    page = re.sub(r"<script.*?</script>|<style.*?</style>", "", page, flags=re.S)
    return [H.unescape(l).strip() for l in re.sub(r"<[^>]+>", "\n", page).split("\n") if l.strip()]


def num(s):
    m = re.search(r"-?\d+(?:[.,]\d+)?", s or "")
    return float(m.group(0).replace(",", ".")) if m else None


def dopo(L, etichetta, start=0):
    for i in range(start, len(L) - 1):
        if L[i] == etichetta:
            return i, L[i + 1]
    return -1, ""


def storia(L):
    """Paragrafi descrittivi (lunghi) della scheda modello, prima delle specifiche tecniche."""
    fine = next((i for i, l in enumerate(L) if l == "Specifiche tecniche"), len(L))
    scarta = ("brochure", "consent", "consenso", "privacy", "contattaci", "contact us")
    par = [l[0].upper() + l[1:] for l in L[:fine] if len(l) > 90 and not any(k in l.lower() for k in scarta)]
    return "\n\n".join(dict.fromkeys(par[:4]))


def elenco_modelli():
    visti, out = set(), []
    for slug, nome_serie in SERIE.items():
        pag = get(f"{BASE}/it-it/{slug}")
        for m in re.finditer(r'href="(?:https:)?(?://www\.ferretti-yachts\.com)?(/it-it/%s/Model/p/([^"#/]+)/n/([^"#]+))"' % slug, pag):
            path, pid, nome = m.groups()
            if pid in visti or nome in visti:   # varianti -EXT: stessa barca
                continue
            visti.add(pid); visti.add(nome)
            out.append((nome_serie, BASE + path, pid, nome))
        time.sleep(1)
    return out


def parse_modello(serie, url, pid, slug_nome, delay):
    pag = get(url)
    L = righe(pag)
    nome = H.unescape(slug_nome).replace("-", " ")
    i_ini = next((i for i, l in enumerate(L) if l == "Caratteristiche principali"), -1)
    i_mot = next((i for i, l in enumerate(L) if l == "Motorizzazione e velocità"), -1)
    i_tec = next((i for i, l in enumerate(L) if l == "Specifiche tecniche"), -1)
    # descrizione: primo paragrafo lungo dopo il nome
    desc = next((l for l in L[:120] if len(l) > 90), "")
    r = {c: "" for c in COLONNE}
    r.update(Codice=pid.replace("-PUB", ""), Famiglia=serie, Categoria="Motor yacht", Sottocategoria=serie,
             Nome=nome, Prezzo=0, Descrizione=desc, URL=url)
    for etichetta, col in (("LOA", "LOA_m"), ("Larghezza massima", "Larghezza_m"), ("Pescaggio", "Pescaggio_m")):
        _, v = dopo(L, etichetta, max(i_ini, 0))
        r[col] = num(v) or ""
    _, v = dopo(L, "LH", max(i_tec, 0))
    r["Lunghezza_scafo_m"] = num(v) or ""
    for etichetta, col in (("Dislocamento a vuoto", "Dislocamento_Vuoto_kg"), ("Dislocamento a pieno carico", "Dislocamento_Carico_kg"),
                           ("Serbatoi carburante", "Carburante_l"), ("Serbatoi acqua", "Acqua_l"), ("Nr. persone a bordo", "Persone")):
        _, v = dopo(L, etichetta, max(i_tec, 0))
        r[col] = num(v) or ""
    # motorizzazioni: terne [cabine, motore, Standard|Optional] poi coppie velocita'
    blocco = L[i_mot + 1:i_tec] if 0 <= i_mot < i_tec else []
    motori, cabine, vel = [], "", []
    for i, l in enumerate(blocco):
        if l in ("Standard", "Optional") and i >= 1:
            motori.append((blocco[i - 1], l))
            if i >= 2 and "cabin" in blocco[i - 2].lower():
                cabine = cabine or blocco[i - 2]
    k = 0
    while k < len(blocco) - 1:
        if blocco[k] == "Velocità massima" and k + 3 < len(blocco) and blocco[k + 2] == "Velocità di crociera":
            vel.append((num(blocco[k + 1]), num(blocco[k + 3])))
            k += 4
        else:
            k += 1
    r["Cabine"] = cabine
    r["Storia_IT"] = storia(L)
    for idx, (mot, tipo) in enumerate(motori):
        v = vel[idx] if idx < len(vel) else (None, None)
        suf = "Standard" if tipo == "Standard" else "Optional"
        if r[f"Motore_{suf}"]:
            continue
        r[f"Motore_{suf}"] = mot
        r[f"Vel_Max_{suf}_kn"] = v[0] or ""
        r[f"Vel_Crociera_{suf}_kn"] = v[1] or ""
    # immagini: prima della galleria
    imgs = re.findall(r"https://frg-fwm\.azurewebsites\.net/asset/fwm/Upload/Models/Images/(?:PhotoGallery|Banner)/Ferretti/\d+\.webp", pag)
    uniche = list(dict.fromkeys(imgs))
    r["Immagine_URL"] = uniche[0] if uniche else ""
    try:
        time.sleep(delay)
        Le = righe(get(url.replace("/it-it/", "/en-us/")))
        r["Storia_EN"] = storia(Le)
        r["Descrizione_EN"] = next((l for l in Le[:120] if len(l) > 90), "")
    except Exception as e:
        print("EN KO", url, e, file=sys.stderr)
    time.sleep(delay)
    return r, uniche


def scrivi_excel(rows, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Macchine"
    ws["A1"] = "Catalogo Ferretti Yachts (estratto dal sito ufficiale)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Prezzi a 0: da compilare. Dati tecnici dal sito, da verificare con il listino ufficiale prima di emettere offerte."
    for c, nome in enumerate(COLONNE, 1):
        cell = ws.cell(row=4, column=c, value=nome)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B2447")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = 22 if nome not in ("Descrizione", "Descrizione_EN", "Storia_IT", "Storia_EN", "URL", "Immagine_URL") else 40
    for ri, row in enumerate(rows, 5):
        for c, nome in enumerate(COLONNE, 1):
            ws.cell(row=ri, column=c, value=row.get(nome, ""))
    ws.freeze_panes = "F5"
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="../excel/catalogo_ferretti.xlsx")
    ap.add_argument("--delay", type=float, default=1.5)
    ap.add_argument("--images", default="../images")
    a = ap.parse_args()
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    os.makedirs(a.images, exist_ok=True)
    modelli = elenco_modelli()
    print(f"{len(modelli)} modelli trovati")
    rows = []
    for serie, url, pid, nome in modelli:
        try:
            r, imgs = parse_modello(serie, url, pid, nome, a.delay)
        except Exception as e:
            print("ERRORE", url, e, file=sys.stderr)
            continue
        if r["Immagine_URL"]:
            f = f"{r['Codice']}.webp"
            try:
                with open(os.path.join(a.images, f), "wb") as fh:
                    fh.write(get(r["Immagine_URL"], binary=True))
                r["Immagine_File"] = f
            except Exception as e:
                print("img KO", r["Codice"], e, file=sys.stderr)
        rows.append(r)
        print(r["Codice"], r["Nome"], r["LOA_m"], r["Motore_Standard"], r["Cabine"])
    scrivi_excel(rows, a.out)
    print("scritto", a.out, len(rows))


if __name__ == "__main__":
    main()
