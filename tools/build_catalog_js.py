#!/usr/bin/env python3
"""Converte il catalogo Excel (foglio Macchine) in js/data/catalogo.js, caricato di default dal configuratore.
Uso: python3 tools/build_catalog_js.py excel/catalogo_ferretti.xlsx
Stessa mappatura colonne->campi di js/app/xlsx-loader.js (handleXlsx).
Genera anche gli optional di partenza (motorizzazione alternativa del sito, prezzo 0)."""
import json, re, sys
from pathlib import Path
from openpyxl import load_workbook

ORDINE_SERIE = {"Gamma flybridge": 1, "Infynito": 2}


def fmt(v, unita=""):
    if v in (None, ""):
        return ""
    s = f"{v:g}".replace(".", ",") if isinstance(v, (int, float)) else str(v).strip()
    if isinstance(v, (int, float)) and v >= 1000:
        s = f"{int(v):,}".replace(",", ".") if float(v).is_integer() else s
    return f"{s} {unita}".strip()


def fascia(loa):
    if not isinstance(loa, (int, float)):
        return "—"
    return "Fino a 20 m" if loa < 20 else ("20–25 m" if loa < 25 else ("25–30 m" if loa < 30 else "Oltre 30 m"))


def velocita(vmax, vcro):
    if vmax in (None, "") and vcro in (None, ""):
        return ""
    return f"{fmt(vmax)} / {fmt(vcro)} kn"


def cabine(raw):
    """'4 cabins + 1 crew' -> '4 + 1' (neutro rispetto alla lingua; l'etichetta spiega ospiti + equipaggio)."""
    solo = re.fullmatch(r"\s*([\d/]+)\s*cabins?\s*", str(raw or ""), re.I)
    if solo:
        return solo.group(1)
    m = re.match(r"\s*([\d/]+)\s*cabins?\s*\+\s*(\d+|crew)\s*(?:crew)?", str(raw or ""), re.I)
    return f"{m.group(1)} + {m.group(2)}" if m else str(raw or "").strip()


def main(xlsx):
    ws = load_workbook(xlsx)["Macchine"]
    tutte = list(ws.iter_rows(values_only=True))
    i = next(i for i, r in enumerate(tutte[:15]) if r and "Codice" in r)   # salta titolo/sottotitolo
    righe = [dict(zip(tutte[i], r)) for r in tutte[i + 1:]]
    out, optionals = [], {}
    for r in righe:
        if not r.get("Codice"):
            continue
        code = str(r["Codice"]).strip()
        serie = str(r.get("Famiglia") or "").strip() or "Altro"
        loa = r.get("LOA_m")
        m = {
            "code": code, "family": serie, "category": serie, "subcategory": fascia(loa), "protection": "",
            "name": str(r.get("Nome") or code).strip(), "desc": str(r.get("Descrizione") or "").strip(), "desc_en": str(r.get("Descrizione_EN") or "").strip(),
            "story": str(r.get("Storia_IT") or "").strip(), "story_en": str(r.get("Storia_EN") or "").strip(),
            "loa": fmt(loa, "m"), "beam": fmt(r.get("Larghezza_m"), "m"), "draft": fmt(r.get("Pescaggio_m"), "m"),
            "hull": fmt(r.get("Lunghezza_scafo_m"), "m"), "cabins": cabine(r.get("Cabine")),
            "engine": str(r.get("Motore_Standard") or "").strip(),
            "speed": velocita(r.get("Vel_Max_Standard_kn"), r.get("Vel_Crociera_Standard_kn")),
            "engine_opt": str(r.get("Motore_Optional") or "").strip(),
            "speed_opt": velocita(r.get("Vel_Max_Optional_kn"), r.get("Vel_Crociera_Optional_kn")),
            "disp_empty": fmt(r.get("Dislocamento_Vuoto_kg"), "kg"), "disp_full": fmt(r.get("Dislocamento_Carico_kg"), "kg"),
            "fuel": fmt(r.get("Carburante_l"), "l"), "water": fmt(r.get("Acqua_l"), "l"), "persons": fmt(r.get("Persone")),
            "url": str(r.get("URL") or ""), "image": str(r.get("Immagine_URL") or ""), "imageFile": str(r.get("Immagine_File") or ""),
            "atex": False, "marking": "", "application": "",
            "price": 0, "catOrder": ORDINE_SERIE.get(serie, 99),
        }
        m["photos"] = sorted(f"brochure/{f.name}" for f in (Path(__file__).resolve().parent.parent / "images" / "brochure").glob(f"{code}-*.jpg"))
        m["tagline"] = "Just like home" if serie.startswith("Gamma flybridge") else ""
        out.append(m)
        if m["engine_opt"]:
            optionals[code] = [{"code": "ENG-" + m["name"].replace("Ferretti Yachts ", "").replace(" ", "-").upper(), "desc": m["engine_opt"] + (f" — {m['speed_opt']}" if m["speed_opt"] else ""), "group": "Motorizzazione", "price": 0, "order": 1}]
    dest = Path(__file__).resolve().parent.parent / "js" / "data" / "catalogo.js"
    dest.write_text("/* GENERATO da tools/build_catalog_js.py — non modificare a mano. Fonte: %s */\n"
                    "const CATALOGO_YACHT = %s;\nconst OPTIONALS_YACHT = %s;\n"
                    % (Path(xlsx).name, json.dumps(out, ensure_ascii=False), json.dumps(optionals, ensure_ascii=False)), encoding="utf8")
    print(f"{len(out)} modelli -> {dest}")


if __name__ == "__main__":
    main(sys.argv[1])
